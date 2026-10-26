from __future__ import annotations

import json
from typing import Any
from django.contrib.auth.decorators import login_required
from django.http import HttpRequest, HttpResponse, JsonResponse
from django.shortcuts import render
from django.views.decorators.cache import never_cache
from django.views.decorators.csrf import csrf_exempt

from .services import (
    get_available_brand_groups,
    get_catalog_and_mappings,
    save_single_mapping,
    batch_save_mappings,
    import_excel_parsed_rows,
)
from .models import ProductCategoryMapping


@never_cache
@login_required
def product_mapping_view(request: HttpRequest) -> HttpResponse:
    """Renders the SKU & Product Category Mapping management page."""
    brand_groups = get_available_brand_groups(request.user)

    context: dict[str, Any] = {
        "brand_groups": brand_groups,
        "default_brand_group": brand_groups[0] if brand_groups else "",
        "is_superuser": getattr(request.user, "is_superuser", False),
    }
    return render(request, "data_management/product_mapping.html", context)


@never_cache
@login_required
def api_mapping_data(request: HttpRequest) -> JsonResponse:
    """Returns JSON payload with paginated SKUs, KPIs, and available categories."""
    brand_group = request.GET.get("brand_group", "").lower().strip()
    status_filter = request.GET.get("status", "all").strip().lower()
    category_filter = request.GET.get("category", "ALL").strip()
    search_query = request.GET.get("search", "").strip()

    try:
        page = int(request.GET.get("page", 1))
        page_size = int(request.GET.get("page_size", 25))
    except (ValueError, TypeError):
        page = 1
        page_size = 25

    sort_col = request.GET.get("sort_col", "sku")
    sort_asc = request.GET.get("sort_dir", "asc").lower() != "desc"

    result = get_catalog_and_mappings(
        brand_group=brand_group,
        user=request.user,
        status_filter=status_filter,
        category_filter=category_filter,
        search_query=search_query,
        page=page,
        page_size=page_size,
        sort_col=sort_col,
        sort_asc=sort_asc,
    )
    return JsonResponse(result)


from django.utils.dateparse import parse_datetime
from django.utils import timezone


def _parse_click_timestamp(ts_str: str | None):
    if not ts_str:
        return None
    try:
        parsed = parse_datetime(str(ts_str))
        if parsed:
            if timezone.is_naive(parsed):
                parsed = timezone.make_aware(parsed, timezone.get_current_timezone())
            else:
                parsed = parsed.astimezone(timezone.get_current_timezone())
            return parsed.replace(microsecond=0)
    except Exception:
        pass
    return None


@csrf_exempt
@login_required
def api_save_mapping(request: HttpRequest) -> JsonResponse:
    """POST endpoint to save or update 1 SKU mapping."""
    if request.method != "POST":
        return JsonResponse({"status": "error", "message": "Method not allowed"}, status=405)

    try:
        data = json.loads(request.body.decode("utf-8"))
    except Exception:
        data = request.POST

    brand_group = data.get("brand_group", "").lower().strip()
    sku_code = data.get("sku_code", "").strip()
    product_category = data.get("product_category", "").strip()

    if not brand_group or not sku_code or not product_category:
        return JsonResponse({"status": "error", "message": "brand_group, sku_code, and product_category are required"}, status=400)

    click_dt = _parse_click_timestamp(data.get("load_timestamp"))
    mapping = save_single_mapping(brand_group, sku_code, product_category, request.user, click_time=click_dt)
    return JsonResponse({
        "status": "success",
        "message": f"SKU {sku_code} successfully mapped to {product_category}",
        "data": {
            "key_id": mapping.key_id,
            "brand_group": mapping.brand_group,
            "sku_code": mapping.sku_code,
            "product_category": mapping.product_category,
            "pic": mapping.pic,
            "last_updated": timezone.localtime(mapping.last_updated).strftime("%Y-%m-%d %H:%M:%S"),
        }
    })


@csrf_exempt
@login_required
def api_batch_mapping(request: HttpRequest) -> JsonResponse:
    """POST endpoint to batch update or clear category for multiple SKUs."""
    if request.method != "POST":
        return JsonResponse({"status": "error", "message": "Method not allowed"}, status=405)

    try:
        data = json.loads(request.body.decode("utf-8"))
    except Exception:
        data = request.POST

    brand_group = data.get("brand_group", "").lower().strip()
    sku_codes = data.get("sku_codes", [])
    product_category = data.get("product_category", "").strip()
    clear = bool(data.get("clear", False))

    if not brand_group or not sku_codes:
        return JsonResponse({"status": "error", "message": "brand_group and sku_codes are required"}, status=400)

    click_dt = _parse_click_timestamp(data.get("load_timestamp"))
    updated_count = batch_save_mappings(
        brand_group=brand_group,
        sku_codes=sku_codes,
        product_category=product_category,
        user=request.user,
        clear=clear,
        click_time=click_dt,
    )

    action_word = "dikosongkan" if clear else f"dipetakan ke {product_category}"
    return JsonResponse({
        "status": "success",
        "message": f"Berhasil {action_word} untuk {updated_count} SKU.",
        "updated_count": updated_count,
    })


@csrf_exempt
@login_required
def api_import_excel(request: HttpRequest) -> JsonResponse:
    """POST endpoint to import parsed rows from Excel using bulk create."""
    if request.method != "POST":
        return JsonResponse({"status": "error", "message": "Method not allowed"}, status=405)

    try:
        data = json.loads(request.body.decode("utf-8"))
    except Exception:
        return JsonResponse({"status": "error", "message": "Invalid JSON body"}, status=400)

    rows = data.get("rows", [])
    if not rows:
        return JsonResponse({"status": "error", "message": "Tidak ada baris data untuk diimport"}, status=400)

    click_dt = _parse_click_timestamp(data.get("load_timestamp"))
    res = import_excel_parsed_rows(rows, request.user, click_time=click_dt)

    dup_msg = ""
    if res.get("duplicates_collapsed", 0) > 0:
        dup_msg = f" ({res['duplicates_collapsed']} baris duplikat SKU disatukan)"

    return JsonResponse({
        "status": "success",
        "message": f"Berhasil menyimpan {res['updated']} SKU ke database{dup_msg}.",
        "imported_count": res["updated"],
        "total_rows": res.get("total_rows", res["updated"]),
        "duplicates_collapsed": res.get("duplicates_collapsed", 0),
    })


@never_cache
@login_required
def api_get_categories(request: HttpRequest) -> JsonResponse:
    """Returns list of distinct categories and SKU count for a brand_group."""
    brand_group = request.GET.get("brand_group", "").lower().strip()
    if not brand_group:
        return JsonResponse({"categories": []})

    qs = ProductCategoryMapping.objects.filter(brand_group=brand_group)
    cats = list(qs.values_list("product_category", flat=True).distinct().order_by("product_category"))

    cat_counts = []
    for c in cats:
        count = qs.filter(product_category=c).count()
        cat_counts.append({"name": c, "count": count})

    return JsonResponse({
        "brand_group": brand_group,
        "categories": cat_counts,
    })


@csrf_exempt
@login_required
def api_delete_category(request: HttpRequest) -> JsonResponse:
    """Deletes / unlinks a category from all SKUs in a brand_group."""
    if request.method != "POST":
        return JsonResponse({"status": "error", "message": "Method not allowed"}, status=405)

    try:
        data = json.loads(request.body.decode("utf-8"))
    except Exception:
        data = request.POST

    brand_group = data.get("brand_group", "").lower().strip()
    category = data.get("category", "").strip()

    if not brand_group or not category:
        return JsonResponse({"status": "error", "message": "brand_group and category required"}, status=400)

    deleted_count, _ = ProductCategoryMapping.objects.filter(
        brand_group=brand_group,
        product_category=category,
    ).delete()

    return JsonResponse({
        "status": "success",
        "message": f"Category '{category}' unlinked from {deleted_count} SKUs.",
        "deleted_count": deleted_count,
    })


@never_cache
@login_required
def api_template_catalog(request: HttpRequest) -> JsonResponse:
    """Returns catalog of SKUs for a brand_group to pre-fill in the Excel template."""
    brand_group = request.GET.get("brand_group", "").lower().strip()
    if not brand_group:
        return JsonResponse({"status": "error", "message": "brand_group is required"}, status=400)

    data = get_catalog_and_mappings(brand_group, request.user, page=1, page_size=10000, status_filter="all")

    rows = []
    for item in data.get("items", []):
        rows.append({
            "sku_code": item["sku"],
            "product_name": item["name"] or "",
            "brand_group": item["brand_group"],
            "product_category": item["category"] or "",
            "is_mapped": item["is_mapped"],
        })

    return JsonResponse({
        "status": "success",
        "brand_group": brand_group,
        "rows": rows,
        "total": len(rows),
    })


@never_cache
@login_required
def api_download_template_excel(request: HttpRequest) -> HttpResponse:
    """
    Generates and downloads a lightweight Excel (.xlsx) template
    where entire Column A, B, and C natively have Text format ('@') down to row 1,048,576.
    """
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    brand_group = request.GET.get("brand_group", "").strip() or "mars"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Product_Category_Template"

    # Set whole column dimensions to TEXT format ('@')
    # In Excel OpenXML specification, this writes <col min="1" max="1" style="1" /> where style 1 has numFmtId 49 (@)
    # This guarantees the ENTIRE Column A, B, and C down to row 1,048,576 natively has Text data type in Excel!
    ws.column_dimensions["A"].number_format = "@"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].number_format = "@"
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].number_format = "@"
    ws.column_dimensions["C"].width = 28

    # Headers
    headers = ["sku_code", "brand_group", "product_category"]
    ws.append(headers)

    # Style header row
    header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="0F172A")
    for col_idx in range(1, 4):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.number_format = "@"

    # Sample rows (pure string values)
    sample_rows = [
        ["1729384538798392851", brand_group, "Contoh Kategori A"],
        ["WHIS-DRY-1.2KG", brand_group, "Cat Dry Food"],
        ["PEDI-WET-130G", brand_group, "Dog Wet Food"],
        ["111617-24", brand_group, "Dog Care & Treat"],
    ]

    for row_data in sample_rows:
        ws.append(row_data)

    # Ensure sample data cells also have explicit '@' text format
    for row in ws.iter_rows(min_row=2, max_row=len(sample_rows) + 1, min_col=1, max_col=3):
        for cell in row:
            cell.number_format = "@"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"template_sku_product_category_{brand_group}.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response

